import copy
from datetime import datetime
import json
from pathlib import Path
import tempfile
import unittest

import numpy as np
from openpyxl import load_workbook
import account_replay as A
import backtest_engine as B
from account_statistics import (ACCOUNT_COLUMNS, ACCOUNT_FIELD, CAPACITY_TIMESTAMP, CAPACITY_DATE,
                                account_row, encode_accounts, capacity_date, capacity_timestamp)
from export_names import distinguished_workbook_path
from portable_rank_export import export_portable_view
from ranking_view import build_view, FRONT_HEADERS, config_fingerprint
import test_v155_complete_export as fixtures


class CapacityTests(unittest.TestCase):
    def simulate(self, sizes, *, initial=1000., maximum=20., candidates=(0,2,4), partial=False):
        close=np.array([100.,200.,100.,100.,100.,100.])
        n=len(close); sx=np.array([1,3,3,5,5,5],dtype=np.int64)
        never=np.full(n,n,dtype=np.int64)
        head=(np.array(candidates,dtype=np.int64),np.ones(len(candidates),dtype=np.int8),
              np.arange(n,dtype=np.int32),sx,sx,close[sx],close[sx])
        tail=(close,close,close,np.array(sizes),np.zeros(n,dtype=np.int16),1)
        if partial:
            result=A.simulate_partial_sizes(*head,never,never,close,close,.5,0,never,never,0.,0.,
                *tail,initial=initial,max_qty=maximum)
        else:
            result=A.simulate_all_sizes(*head,never,never,close,close,0,0.,0.,
                *tail,initial=initial,max_qty=maximum)
        return result

    def test_each_account_records_its_own_first_actual_capped_opening(self):
        for partial in (False,True):
            result=self.simulate([.5,1.,2.],partial=partial)
            np.testing.assert_array_equal(result[18],[-1,2,0])
            self.assertEqual([B.metrics_for_size(B.metrics_dict(result),i)['首次达到开仓上限索引']
                              for i in range(3)],[-1,2,0])

    def test_no_trade_or_end_of_sample_capacity_is_not_a_reached_trade(self):
        self.assertEqual(self.simulate([2.],candidates=())[18][0],-1)
        self.assertEqual(self.simulate([2.],candidates=(5,))[18][0],-1)
        self.assertEqual(self.simulate([1.],maximum=21.)[18][0],-1)
        self.assertEqual(self.simulate([1.],maximum=19.)[18][0],2)

    def test_timestamp_uses_close_and_beijing_day_not_the_signal_open_day(self):
        close_times=np.array([1735747140000,1735747200000],dtype=np.int64) # UTC 15:59 / 16:00
        self.assertEqual(capacity_timestamp(1,close_times),1735747200000)
        self.assertEqual(capacity_date(close_times[1])-capacity_date(close_times[0]),1)
        self.assertEqual(capacity_date(-1),'未达到')
        self.assertEqual(capacity_date(None),'未记录（旧结果）')
        self.assertEqual(capacity_date(-2),'未记录（旧结果）')

    def test_old_packed_statistics_remain_readable_and_unknown_is_not_never(self):
        for length in (12,14):
            restored=account_row({ACCOUNT_FIELD:'|'.join(['1']*length)},0)
            self.assertNotIn(CAPACITY_TIMESTAMP,restored)
        old={key:1. for key in ACCOUNT_COLUMNS if key!=CAPACITY_TIMESTAMP}
        self.assertEqual(float(account_row({ACCOUNT_FIELD:encode_accounts([old])},0)[CAPACITY_TIMESTAMP]),-2)


class ExportReadingTests(unittest.TestCase):
    def setUp(self):
        self.fixture=fixtures.CompleteExportTests();self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)

    def test_date_layout_first_sheet_colors_and_original_fingerprints(self):
        payload=copy.deepcopy(self.fixture.payload)
        payload['表头'].append(CAPACITY_TIMESTAMP)
        for row,stamp in zip(payload['分类']['移动止盈'],[1735747200000,-1]):row.append(stamp)
        view=build_view(payload)
        self.assertEqual(view['表头'][:7],FRONT_HEADERS)
        self.assertEqual(view['表头'][7],CAPACITY_DATE)
        target=self.fixture.root/'readable.xlsx'
        export_portable_view(view,target)
        book=load_workbook(target)
        try:
            self.assertEqual(book.sheetnames[0],'全局前5000')
            self.assertEqual(book.active.title,'全局前5000')
            sheet=book.active
            self.assertEqual([c.value for c in sheet[1]][:7],FRONT_HEADERS)
            self.assertEqual(sheet.freeze_panes,'B2')
            self.assertEqual(sheet['H2'].value,datetime(2025,1,2))
            self.assertEqual(sheet['H3'].value,'未达到')
            self.assertEqual(sheet['A2'].value,config_fingerprint(self.fixture.rows[0]))
            rules={str(c.sqref):sheet.conditional_formatting[c] for c in sheet.conditional_formatting}
            self.assertEqual([c.rgb[-6:] for c in rules['C2:C3'][0].colorScale.color],['63BE7B','FFEB84','F8696B'])
            self.assertEqual(len(rules['G2:G3']),4)
            self.assertEqual(sheet['G2'].fill.fgColor.rgb[-6:],'FCE4D6')
            self.assertEqual(sheet['G3'].fill.fgColor.rgb[-6:],'DDEBF7')
            self.assertEqual(sheet['B2'].value,self.fixture.rows[0]['期末资金（USDC）'])
            self.assertEqual(sheet['I2'].value,self.fixture.rows[0]['ETH单次最大开仓数量（ETH）'])
        finally:book.close()

    def test_legacy_results_display_unrecorded(self):
        view=build_view(self.fixture.payload)
        self.assertTrue(all(row[7]=='未记录（旧结果）' for rows in view['分类'].values() for row in rows))

    def test_distinct_directories_keep_stable_two_digit_suffixes(self):
        root=self.fixture.root
        name='各类止盈最优前5000名.xlsx'
        first=distinguished_workbook_path(root/'one'/name,root)
        second=distinguished_workbook_path(root/'two'/name,root)
        self.assertEqual(first.name,'各类止盈最优前5000名_01.xlsx')
        self.assertEqual(second.name,'各类止盈最优前5000名_02.xlsx')
        self.assertEqual(distinguished_workbook_path(root/'one'/name,root),first)
        self.assertEqual(distinguished_workbook_path(root/'one'/'各类止盈最差1000名.xlsx',root).name,
                         '各类止盈最差1000名_01.xlsx')
        self.assertEqual(distinguished_workbook_path(root/'one'/'各类止盈最优前10000名.xlsx',root).name,
                         '各类止盈最优前10000名_01.xlsx')
