"""Excel must preserve generated 64-bit rule IDs as text identifiers."""
import csv
from pathlib import Path
import tempfile
import unittest

from hedge_config import normalize_config
from hedge_worker import summarize, summary_headers, csv_rows, export_excel, CASE_HEADERS


class CompositeCodeExportTests(unittest.TestCase):
    def test_large_composite_code_and_period_codes_match_csv_exactly(self):
        from openpyxl import load_workbook
        code = -2631998912456684628
        descriptor = dict(id="fixture", label="复合规则精度", timeframe="5m+1m", code=code,
                          cases=(0, 0, 0, code, 42), field="dif", entry_mode="LIVE_01", s3_gap=.001, s3_timeframe="1m")
        stats = dict(ending_equity=23456.789, max_drawdown=.125, profit_factor=1.4, trades=100,
                     win_rate=.6, daily_trades=.5, fees=0., unrealized_pnl=0.)
        summary = summarize(descriptor, normalize_config(), 6., .001, 0, [stats])
        with tempfile.TemporaryDirectory() as folder:
            directory = Path(folder)
            csv_rows(directory / "summary.csv", summary_headers([summary]), [summary])
            with (directory / "summary.csv").open(encoding="utf-8-sig", newline="") as stream:
                csv_row = next(csv.DictReader(stream))
            target = export_excel(directory, [summary], [], [], {})
            book = load_workbook(target, read_only=False, data_only=True)
            page = book["方案汇总"]
            columns = {cell.value: cell.column for cell in page[1]}
            for name in ["开仓代码", *CASE_HEADERS]:
                cell = page.cell(2, columns[name])
                self.assertEqual(cell.data_type, "s", name)
                self.assertEqual(cell.value, csv_row[name], name)
            self.assertEqual(page.cell(2, columns["开仓代码"]).value, str(code))
            self.assertEqual(page.cell(2, columns["期末资金中位数（USDC）"]).value, 23456.789)
            self.assertEqual(page.cell(2, columns["期末资金中位数（USDC）"]).data_type, "n")
            book.close()


if __name__ == "__main__":
    unittest.main()
