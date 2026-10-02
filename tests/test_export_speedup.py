"""Export validation reuse keeps independent complete configurations and workbook semantics."""
import copy
import json
from pathlib import Path
import unittest
from unittest import mock

import export_runtime
import fingerprint_lookup
from ranking_view import CODE_NAMES, build_view
from portable_rank_export import export_portable_view
import strategy_description as description
import test_v155_complete_export as fixture


class ExportSpeedupTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixture.CompleteExportTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)

    def test_narrowed_configurations_do_not_mutate_or_share_the_original_context(self):
        payload = copy.deepcopy(self.fixture.payload)
        with mock.patch.object(description, '生成止损组合', wraps=description.生成止损组合) as stops:
            context = description.load_context(payload, None, CODE_NAMES)
            before = copy.deepcopy(context)
            first = description._single_selection(self.fixture.rows[0], context)
            second = description._single_selection(self.fixture.rows[1], context)
            self.assertEqual(stops.call_count, 1, 'The stop catalogue must be read once per export context')
        self.assertEqual(context, before)
        self.assertEqual(payload, self.fixture.payload)
        self.assertEqual(first['成本模式'], 'SLIPPAGE')
        self.assertEqual(second['成本模式'], 'FEE')
        first['资金约束']['初始资金USDC'] = 999
        first['开仓条件']['1m'].append(0)
        first['候选筛选']['启用'] = not first['候选筛选']['启用']
        self.assertEqual(context, before)
        self.assertEqual(second['资金约束'], before['selection']['资金约束'])
        self.assertEqual(second['开仓条件']['1m'], [152])

    def test_reused_validation_still_rejects_dimension_fee_mode_and_identity_conflicts(self):
        context = description.load_context(self.fixture.payload, None, CODE_NAMES)
        invalid = {header: None for header in fingerprint_lookup.SINGLE_FIELDS.values()}
        invalid.update({'基础策略编号': self.fixture.rows[0]['基础策略编号'] + 1,
                        '开仓MACD口径': 'invalid', '1分钟条件': '不存在的条件',
                        '开仓位置过滤代码': 'invalid', '入场触发口径': 'invalid', '成本模式': 'invalid',
                        '开仓净手续费率（%）': .99, '开仓成交偏移（%）': .99,
                        '回测计算版本': 'invalid'})
        for key, value in invalid.items():
            with self.subTest(field=key):
                row = dict(self.fixture.rows[0], **{key: value})
                result = description.describe_row(row, context)
                self.assertEqual(result['完整单策略配置JSON'], '')
                self.assertIn('不完整', result['参数完整性'])

    def test_rows_do_not_rebuild_catalogues_after_context_validation(self):
        import selection_config
        context = description.load_context(self.fixture.payload, None, CODE_NAMES)
        with mock.patch.object(selection_config, '规范化配置',
                               side_effect=AssertionError('Already validated')):
            for original in self.fixture.rows:
                selection = description._single_selection(original, context)
                self.assertEqual(selection['成本模式'], original['成本模式'])
                self.assertEqual(selection['入场触发口径'], original['入场触发口径'])

    def test_remaining_extra_choices_are_rejected_by_single_combination_assertion(self):
        context = description.load_context(self.fixture.payload, None, CODE_NAMES)
        narrow = fingerprint_lookup._narrow_strategy_dimensions
        def not_fully_narrowed(*args, **kwargs):
            selection, cases, stop = narrow(*args, **kwargs)
            selection['仓位倍数'] = [1.0, 2.0]
            return selection, cases, stop
        with mock.patch.object(fingerprint_lookup, '_narrow_strategy_dimensions',
                               side_effect=not_fully_narrowed):
            with self.assertRaisesRegex(ValueError, '不是单一完整组合'):
                description._single_selection(self.fixture.rows[0], context)

    def test_compound_entry_stop_and_tp_still_narrow_to_one_complete_strategy(self):
        import test_indicator_combination_export_restore as combinations
        helper = combinations.CombinationExportRestoreTests()
        helper.setUp()
        self.addCleanup(helper.doCleanups)
        directory = helper.result()
        context = description.load_context(helper.payload, directory, CODE_NAMES)
        import selection_config
        with mock.patch.object(selection_config, '规范化配置',
                               side_effect=AssertionError('Already validated')):
            selection = description._single_selection(helper.rows[0], context)
        self.assertEqual(selection['开仓条件']['1m'], [5, 152])
        self.assertEqual(selection['开仓条件']['5m'], [5, 152])
        self.assertEqual(selection['止损代码'], ['S1', 'S3_1m'])
        self.assertEqual(selection['止盈方案编号'], [2, 8427])
        self.assertEqual(selection_config.配置统计(selection)['包含仓位完整组合数'], 1)

    def test_shared_cell_styles_keep_number_formats_text_ids_formulas_and_empty_sheet_layout(self):
        from openpyxl import load_workbook
        view = build_view(self.fixture.payload)
        view['分类'] = {'空表': [], **view['分类']}
        output = self.fixture.root / 'styles.xlsx'
        export_portable_view(view, output)
        book = load_workbook(output, data_only=False)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames[:3], ['全局前5000', '空表', '移动止盈'])
        empty = book['空表']
        self.assertEqual(empty.freeze_panes, 'B2')
        self.assertEqual(empty.max_row, 1)
        self.assertTrue(empty.auto_filter.ref.endswith('1'))
        sheet = book['移动止盈']
        columns = {cell.value: cell.column for cell in sheet[1]}
        for number, raw in enumerate(self.fixture.rows, 2):
            self.assertEqual(sheet.cell(number, columns['基础策略编号']).data_type, 's')
            self.assertEqual(sheet.cell(number, columns['基础策略编号']).number_format, '@')
            self.assertEqual(sheet.cell(number, columns['完整单策略配置JSON']).data_type, 's')
            self.assertEqual(sheet.cell(number, columns['账户净利润（USDC）']).data_type, 'f')
            self.assertEqual(sheet.cell(number, columns['开仓成交偏移（%）']).number_format,
                             '0.000000%;[Red]-0.000000%')
            self.assertEqual(sheet.cell(number, columns['期末资金（USDC）']).number_format,
                             '#,##0.00;[Red]-#,##0.00')
            for cell in sheet[number]:
                self.assertEqual(cell.alignment.horizontal, 'center')
                self.assertTrue(cell.alignment.wrap_text)
            self.assertEqual(sheet.row_dimensions[number].height, 40)
        self.assertEqual(len(sheet.conditional_formatting), 6)

    def test_locked_final_workbook_keeps_original_and_saves_complete_new_file(self):
        source = self.fixture.root / 'ranking.json'
        source.write_text(json.dumps(self.fixture.payload, ensure_ascii=False), 'utf-8')
        output = self.fixture.root / 'opened.xlsx'
        output.write_bytes(b'original opened workbook')
        replace = export_runtime.os.replace
        def locked_once(temporary, target):
            if Path(target) == output:
                raise PermissionError('Excel holds the existing workbook')
            return replace(temporary, target)
        with mock.patch.dict('sys.modules', {'artifact_tool': None}), \
             mock.patch.object(export_runtime.os, 'replace', side_effect=locked_once):
            actual = export_runtime.export_xlsx_atomic(None, Path('excel_export.mjs'), source,
                                                       output, source.parent)
        self.assertEqual(output.read_bytes(), b'original opened workbook')
        self.assertNotEqual(actual, output)
        self.assertTrue(actual.name.startswith('opened_本次_'))
        from openpyxl import load_workbook
        book = load_workbook(actual, read_only=True)
        try:
            self.assertIn('移动止盈', book.sheetnames)
        finally:
            book.close()
        self.assertEqual(list(self.fixture.root.glob('*.tmp.xlsx')), [])


if __name__ == '__main__':
    unittest.main()
