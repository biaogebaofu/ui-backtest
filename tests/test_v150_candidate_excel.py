import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from streaming_excel_export import export_streaming_xlsx


class CandidateWorkbookTests(unittest.TestCase):
    def test_fallback_retains_candidates_and_stage_counts(self):
        headers = ['策略指纹', '期末资金（USDC）', '最大回撤（%）', '名义倍数（倍）',
                   '胜率（%）', '平均日完整交易数（次/日）']
        candidate = dict(zip(headers, ['0123456789abcdef', 1900., .05, 2., .6, 7.5]))
        payload = {'表头': headers, '研究候选': [candidate], '观察候选': [candidate], '实盘候选': [],
                   '设置': {'筛选方案': 'RETURN_DRAWDOWN', '比较范围': 'ALL_RUN', '资金保留比例': .9},
                   '总行数': 3, '初筛通过行数': 2, '研究候选数': 1, '观察候选数': 1,
                   '实盘候选数': 0, '筛选阶段统计': {'资金保留门槛通过': 1},
                   '研究预筛分说明': '合格最高资金的90%以内按低回撤优先',
                   '淘汰统计': [{'淘汰原因': '资金不足', '涉及行数': 1}],
                   '高级待验证项': ['样本外验证'], '源文件': 'test.csv'}
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source, output = root / 'payload.json', root / 'out.xlsx'
            source.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
            export_streaming_xlsx(source, output)
            book = load_workbook(output)
            try:
                self.assertIn('全局观察候选', book.sheetnames)
                sheet = book['全局观察候选']
                display_headers = [headers[0], '入场次数规则', '成本模式'] + headers[1:]
                self.assertEqual(list(next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))), display_headers)
                displayed = dict(zip(display_headers,next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))))
                self.assertEqual({h:displayed[h] for h in headers},candidate)
                self.assertEqual(displayed['入场次数规则'],'未记录（查原结果）')
                self.assertEqual(len(sheet.conditional_formatting), 6)
                self.assertEqual(sheet.freeze_panes, 'A5')
                self.assertEqual(sheet.auto_filter.ref, 'A4:H5')
                self.assertNotIn('每日收益测算', book.sheetnames)
                guide = '\n'.join(str(c.value) for row in book['0_说明与检查'] for c in row if c.value is not None)
                self.assertIn('资金保留门槛通过', guide)
                self.assertIn('样本外验证', guide)
                self.assertIn('全部已回测倍数', guide)
                self.assertTrue(any('尚未通过实盘验证' in str(c.value) for row in book['实盘候选'] for c in row))
            finally:
                book.close()

    def test_candidate_javascript_guide_has_mode_specific_scope(self):
        source = (Path(__file__).resolve().parents[1] / 'candidate_excel_export.mjs').read_text('utf-8')
        self.assertIn('RETURN_DRAWDOWN', source)
        self.assertIn('ALL_RUN', source)
        self.assertIn('资金保留比例', source)


if __name__ == '__main__':
    unittest.main()
