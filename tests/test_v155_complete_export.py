"""All strategy settings remain screenshot/copy accessible without guessed context."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest import mock

from account_statistics import ENGINE_VERSION
from execution_settings import effective_fee_rates, effective_slippage
from extended_rules import stable_base_id
from ranking_view import build_view, config_fingerprint, CODE_NAMES
from selection_config import 规范化配置, 配置签名, 配置统计
from strategy_space import 生成止损组合
import strategy_description as D
import test_v149_fingerprint_lookup as fixture


class CompleteExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='v155_complete_export_')
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.helper = fixture.FingerprintLookupTests(); self.helper.root = self.root
        self.directory, _, selected = self.helper.make_result('original')
        old = json.loads((self.directory / '各类止盈前5000名.json').read_text('utf-8'))
        template = dict(zip(old['表头'], next(iter(old['分类'].values()))[0]))
        selected.update({'开仓指标': ['hist'], '开仓条件': {'4h':[0], '1h':[0], '15m':[0], '5m':[0], '1m':[5,152]},
                         '止损代码':['ATR150_5m'], '止盈方案编号':[8427], '开仓位置过滤':['OFF'],
                         '成本模式':['SLIPPAGE','FEE'], '入场触发口径':['LIVE_01','TF_EVENT','MACD_CYCLE']})
        self.selection = 规范化配置(selected)
        self.rows = []
        stop_index = next(i for i,x in enumerate(生成止损组合()) if x[0]=='ATR150_5m')
        for cost in ('SLIPPAGE', 'FEE'):
            config = dict(self.selection, 成本模式=cost)
            fees, slip = effective_fee_rates(config), effective_slippage(config)
            row = dict(template, **{'基础策略编号':stable_base_id(0,(0,0,0,0,152),stop_index),
                '开仓MACD口径':'MACD柱', '1分钟条件':CODE_NAMES[152], '止损代码':'ATR150_5m',
                '止损说明':'ATR止损 5m 1.5倍ATR14', '开仓位置过滤代码':'OFF', '开仓位置过滤说明':'关闭位置过滤（原策略基线）',
                '止盈方案编号':8427,'止盈类别':'移动止盈','止盈周期组合':'','止盈指标':'最高浮盈回吐比例',
                '止盈参数一':.0005,'止盈参数二':.025,'止盈参数三':0., '成本模式':cost,
                '入场触发口径':'TF_EVENT', '开仓成交偏移（%）':slip[0], '平仓成交偏移（%）':slip[1],
                '往返成交偏移（%）':sum(slip), '开仓净手续费率（%）':fees[0], '平仓净手续费率（%）':fees[1],
                '期末资金（USDC）':10001.125, '最大回撤（%）':.12, '胜率（%）':.8, '交易次数（单）':150,
                '平均日完整交易数（次/日）':10., '回测计算版本':ENGINE_VERSION+':'+config['成交价格口径']})
            self.rows.append(row)
        data = json.loads((self.directory/'回测数据说明.json').read_text('utf-8'))
        data['macd'] = '12/26/9，MACD柱=2×(DIF-DEA)'
        identity = json.loads((self.directory/'回测运行身份.json').read_text('utf-8'))
        identity['selection_signature'] = 配置签名(self.selection)
        self.context = {'schema':1,'selection':self.selection,'data':data,'identity':identity}
        self.payload = {'表头':list(self.rows[0]),'分类':{'移动止盈':[list(row.values()) for row in self.rows]},
                        '运行上下文':self.context}

    def view_rows(self, view):
        return [dict(zip(view['表头'], row)) for row in view['分类']['移动止盈']]

    def test_both_costs_scalar_exact_config_complete_settings_and_original_values(self):
        view = build_view(self.payload)
        rows = self.view_rows(view)
        self.assertEqual(len(rows),2)
        for original, row in zip(self.rows,rows):
            envelope = json.loads(row['完整单策略配置JSON'])
            config = envelope['selection']
            self.assertEqual(config['成本模式'],original['成本模式'])
            self.assertEqual(config['入场触发口径'],'TF_EVENT')
            self.assertEqual(配置统计(config)['包含仓位完整组合数'],1)
            self.assertEqual(config['仓位倍数'],[original['名义倍数（倍）']])
            self.assertEqual(config['资金约束'],self.selection['资金约束'])
            self.assertEqual(config['入场约束'],self.selection['入场约束'])
            self.assertEqual(config['成交偏移'],self.selection['成交偏移'])
            self.assertEqual(config['手续费'],self.selection['手续费'])
            self.assertEqual(envelope['sources'],self.context['data']['sources'])
            self.assertEqual(envelope['request'],self.context['data']['request'])
            self.assertEqual(row['策略指纹'],config_fingerprint(original))
            self.assertEqual(row['期末资金（USDC）'],original['期末资金（USDC）'])
            self.assertEqual(row['配置开仓偏移（%）'],.0003)
            self.assertEqual(row['开仓成交偏移（%）'],original['开仓成交偏移（%）'])
            self.assertEqual(row['基础策略编号'],str(original['基础策略编号']))
            self.assertIn('未记录',row['数据标的（元数据）'])
            self.assertIn('真实参数=[9]',row['开仓判据（当前代码解释）'])
            self.assertIn('WMA',row['开仓判据（当前代码解释）'])
            self.assertIn('1.5×最新已收5m',row['止损判据（当前代码解释）'])
            self.assertIn('激活当根不按新线退出',row['止盈判据（当前代码解释）'])
            self.assertIn('不匹配',row['定义来源与版本核对'])
            self.assertIn('不能只凭截图复写',row['参数完整性'])
        self.assertEqual(rows[1]['开仓成交偏移（%）'],0.)
        self.assertIn('成本模式',view['表头'])
        self.assertIn('保护止损浮亏比例（%）',view['表头'])

    def test_partial_take_profit_does_not_reuse_partial_fraction_as_retracement(self):
        from strategy_space import 生成止盈方案
        tp = next(tp for tp in 生成止盈方案() if tp.编号 == 2607)
        row = dict(self.rows[0], **{'止盈方案编号':tp.编号, '止盈类别':tp.类别, '止盈指标':tp.指标,
                    '止盈参数一':tp.参数一, '止盈参数二':tp.参数二, '止盈参数三':tp.参数三})
        context = D.load_context(self.payload, None, CODE_NAMES)
        rule = D.describe_row(row, context)['止盈判据（当前代码解释）']
        self.assertIn('参数2为首段平仓比例，不是回吐比例',rule)
        self.assertIn('余仓止盈方案编号',rule)
        self.assertNotIn('保留(1−参数2)',rule)

    def test_legacy_directory_context_is_read_once_and_missing_or_changed_context_not_guessed(self):
        raw = copy.deepcopy(self.context)
        self.helper.write_json(self.directory/'组合选择.json',raw['selection'])
        self.helper.write_json(self.directory/'回测运行身份.json',raw['identity'])
        self.helper.write_json(self.directory/'回测数据说明.json',raw['data'])
        self.helper.write_json(self.directory/'断点记录.json',{'selection':raw['selection'], 'run_identity':raw['identity']})
        payload = {k:v for k,v in self.payload.items() if k!='运行上下文'}
        with mock.patch.object(D,'_read',wraps=D._read) as read:
            view = build_view(payload,self.directory)
        self.assertEqual(len([c for c in read.call_args_list if c.args[0].name=='组合选择.json']),1)
        self.assertTrue(all(row['完整单策略配置JSON'] for row in self.view_rows(view)))
        self.assertTrue(all(not row['完整单策略配置JSON'] for row in self.view_rows(build_view(payload))))
        changed = copy.deepcopy(self.payload)
        changed['运行上下文']['selection']['资金约束']['初始资金USDC'] += 1
        broken = self.view_rows(build_view(changed))
        self.assertTrue(all(not row['完整单策略配置JSON'] for row in broken))
        self.assertTrue(all('不完整' in row['参数完整性'] for row in broken))

    def test_scalar_cost_remains_visible_when_every_row_is_constant(self):
        payload = copy.deepcopy(self.payload)
        payload['分类']['移动止盈'] = payload['分类']['移动止盈'][:1]
        view = build_view(payload)
        row = self.view_rows(view)[0]
        for key in D.PARAMETER_HEADERS + D.EXTRA_HEADERS:
            # Existing human-readable aliases stay left; their native source is not duplicated.
            self.assertTrue(key in view['表头'] or key in ('止盈后等待分钟',))
        self.assertEqual(row['成本模式代码'],'SLIPPAGE')
        self.assertEqual(row['成本模式'],'仅成交偏移')
        self.assertEqual(row['开仓方向'],'LONG')
        self.assertEqual(row['保护止损浮亏比例（%）'],.73)

    def test_portable_workbook_preserves_numbers_ids_formulas_and_bounded_layout(self):
        from openpyxl import load_workbook
        from portable_rank_export import export_portable_view
        view = build_view(self.payload)
        target = self.root/'temporary_export.xlsx'
        export_portable_view(view,target)
        wb = load_workbook(target,data_only=False)
        self.addCleanup(wb.close)
        sheet = wb['移动止盈']; headers = {cell.value:cell.column for cell in sheet[1]}
        for offset, original in enumerate(self.rows,2):
            fee = sheet.cell(offset,headers['开仓成交偏移（%）'])
            self.assertIsInstance(fee.value,(int,float))
            self.assertEqual(fee.value,original['开仓成交偏移（%）'])
            self.assertIn('0.000000%',fee.number_format)
            self.assertEqual(sheet.cell(offset,headers['基础策略编号']).value,str(original['基础策略编号']))
            self.assertEqual(sheet.cell(offset,headers['基础策略编号']).number_format,'@')
            self.assertGreaterEqual(sheet.column_dimensions[sheet.cell(1,headers['基础策略编号']).column_letter].width,22)
            self.assertTrue(sheet.cell(offset,headers['账户净利润（USDC）']).value.startswith('='))
            self.assertLessEqual(sheet.row_dimensions[offset].height,60)
            self.assertEqual(json.loads(sheet.cell(offset,headers['完整单策略配置JSON']).value)['selection']['成本模式'],original['成本模式'])
            self.assertEqual(len(sheet.conditional_formatting),6)

    def test_candidate_fallback_keeps_extended_columns_precise_and_bounded(self):
        from openpyxl import load_workbook
        from streaming_excel_export import export_streaming_xlsx
        rows = self.view_rows(build_view(self.payload))
        headers = ['策略指纹','期末资金（USDC）','最大回撤（%）','名义倍数（倍）','胜率（%）',
                   '日均完整交易（次/日）','基础策略编号','配置开仓偏移（%）',
                   '开仓判据（当前代码解释）','完整单策略配置JSON']
        source, target = self.root/'candidate.json', self.root/'candidate.xlsx'
        payload = {'表头':headers,'研究候选':rows,'观察候选':rows,'实盘候选':[],'设置':{}}
        self.helper.write_json(source,payload)
        export_streaming_xlsx(source,target)
        book = load_workbook(target); self.addCleanup(book.close)
        sheet = book['每类研究候选']
        actual_headers = [cell.value for cell in sheet[4]]
        expected_headers = [headers[0], '入场次数规则', '成本模式'] + headers[1:]
        self.assertEqual(actual_headers,expected_headers)
        for index, original in enumerate(rows,5):
            self.assertEqual(sheet.cell(index,9).value,original['基础策略编号'])
            self.assertEqual(sheet.cell(index,9).number_format,'@')
            self.assertEqual(sheet.cell(index,10).value,original['配置开仓偏移（%）'])
            self.assertEqual(sheet.cell(index,10).number_format,'0.000000%')
            self.assertEqual(sheet.cell(index,12).value,original['完整单策略配置JSON'])
            self.assertEqual(sheet.row_dimensions[index].height,44)
        self.assertEqual(sheet.column_dimensions['I'].width,24)
        self.assertEqual(sheet.column_dimensions['K'].width,55)
        self.assertEqual(sheet.column_dimensions['L'].width,60)
        self.assertEqual(len(sheet.conditional_formatting),6)
        payload['研究候选'][0]['完整单策略配置JSON'] = 'x'*32768
        self.helper.write_json(source,payload)
        with self.assertRaisesRegex(ValueError,'32767.*第5行.*拒绝截断'):
            export_streaming_xlsx(source,self.root/'oversized_candidate.xlsx')
        self.assertFalse((self.root/'oversized_candidate.xlsx').exists())

    def test_candidate_javascript_has_matching_precision_text_and_length_guards(self):
        source = (Path(__file__).resolve().parents[1]/'candidate_excel_export.mjs').read_text('utf-8')
        self.assertLess(source.index('超过32767字符'),source.index('Workbook.create()'))
        self.assertIn('0.000000%',source)
        self.assertIn('range.format.numberFormat = "@"',source)
        self.assertIn('String(row[header])',source)
        self.assertIn('rowHeight: 44',source)

    def test_oversized_configuration_fails_explicitly_before_either_exporter(self):
        from portable_rank_export import export_portable_view
        from artifact_rank_export import export_artifact_view
        view = {'表头':['完整单策略配置JSON'],'分类':{'主榜':[['x'*32768]]}}
        for exporter in (export_portable_view,export_artifact_view):
            target = self.root/(exporter.__name__+'.xlsx')
            with self.assertRaisesRegex(ValueError,'32767.*拒绝截断'):
                exporter(view,target)
            self.assertFalse(target.exists())


if __name__=='__main__':
    unittest.main()
