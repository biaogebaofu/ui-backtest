"""Real worker, fingerprint restore and Chinese exports for complete MACD cycles."""
import copy
import json
from pathlib import Path
import tempfile
import unittest

from fingerprint_lookup import find_fingerprint_matches, restore_fingerprint_match
from portable_rank_export import export_portable_view
from ranking_view import build_view, config_fingerprint
from selection_config import 规范化配置, 配置统计
from strategy_display import mode_label
import test_v155_multi_cost_backend as worker_fixture


MODES = ['MACD_FULL_RED', 'MACD_FULL_GREEN']


class FullMacdCycleExportTests(unittest.TestCase):
    # Reuse the existing bounded worker harness, not its larger test matrix.
    run_worker = classmethod(worker_fixture.MultiCostRealWorkerTests.run_worker.__func__)

    @classmethod
    def setUpClass(cls):
        import numpy as np
        import pandas as pd

        cls.temp = tempfile.TemporaryDirectory(prefix='full_macd_cycle_export_')
        cls.addClassCleanup(cls.temp.cleanup)
        cls.root = Path(cls.temp.name)
        rng = np.random.default_rng(20260923)
        count = 1800
        close = 2000 * np.exp(np.cumsum(rng.normal(0, .001, count)))
        opening = np.r_[close[0], close[:-1]]
        volume = rng.lognormal(6, 1, count)
        cls.data = cls.root / 'synthetic.csv'
        pd.DataFrame({'openTime': 1735689600000 + np.arange(count) * 60000,
                      'open': opening, 'high': np.maximum(opening, close) + .5,
                      'low': np.minimum(opening, close) - .5,
                      'close': close, 'volume': volume}).to_csv(cls.data, index=False)
        config = worker_fixture.small_config()
        config['开仓条件']['1m'] = [3]
        config.update({'入场触发口径': MODES, '成本模式': 'FEE',
                       '开仓位置过滤': ['OFF'], '止盈方案编号': [2],
                       '强制时间止损分钟': [5], '仓位倍数': [1.],
                       '手续费': {'开仓费率': 0., '平仓费率': 0., 'BNB抵扣': False, '返佣比例': 0.}})
        config['入场约束']['最小S3距离'] = 0.
        config['候选筛选'].update(启用=False, 自动导出=False, 导出旧排行=True)
        cls.config = 规范化配置(config)
        cls.runs = {'combined': cls.run_worker('combined', cls.config)}
        for mode in MODES:
            cls.runs[mode] = cls.run_worker(mode, dict(cls.config, 入场触发口径=mode))
        cls.runs['MACD_CYCLE'] = cls.run_worker('MACD_CYCLE', dict(cls.config, 入场触发口径='MACD_CYCLE'))

    def test_combined_is_exactly_two_independent_anchored_runs(self):
        combined = self.runs['combined']
        self.assertEqual(len(combined['rows']), 2)
        self.assertEqual(len(combined['ranking']), 2)
        fingerprints = {config_fingerprint(row) for row in combined['ranking']}
        self.assertEqual(len(fingerprints), 2)
        self.assertEqual(配置统计(self.config)['包含仓位完整组合数'], 2)
        for mode in MODES:
            self.assertEqual([row for row in combined['rows'] if row['入场触发口径'] == mode],
                             self.runs[mode]['rows'])
            self.assertEqual([row for row in combined['ranking'] if row['入场触发口径'] == mode],
                             self.runs[mode]['ranking'])
            self.assertGreater(int(self.runs[mode]['rows'][0]['交易次数（单）']), 0)
            self.assertLess(int(self.runs[mode]['rows'][0]['交易次数（单）']),
                            int(self.runs['MACD_CYCLE']['rows'][0]['交易次数（单）']))
        stats = ['交易次数（单）', '所选仓位期末资金（USDC）', '平均单笔收益率（%）']
        self.assertNotEqual([self.runs[MODES[0]]['rows'][0][key] for key in stats],
                            [self.runs[MODES[1]]['rows'][0][key] for key in stats])
        start = next(event for event in combined['events'] if event['type'] == 'start')
        self.assertEqual((start['csv_rows'], start['total_combinations']), (2, 2))

    def test_worker_announces_full_cycle_not_each_color_change(self):
        run = self.runs['combined']
        stages = [event.get('message', '') for event in run['events'] if event['type'] == 'stage']
        for mode, boundary in zip(MODES, ('绿转红', '红转绿')):
            note = next(text for text in stages if f'{mode}｜' in text)
            self.assertIn(boundary, note)
            self.assertIn('初始截断色段不开仓', note)
            self.assertIn('多空共用一次实际开仓', note)
            self.assertIn('平仓均不恢复次数', note)
        notes = (run['output'] / '中文字段说明.txt').read_text('utf-8')
        self.assertIn('MACD_CYCLE始终按1m hist', notes)
        self.assertIn('MACD_FULL_RED固定从绿转红边界', notes)
        self.assertIn('MACD_FULL_GREEN固定从红转绿边界', notes)
        self.assertIn('不能在同一策略内叠加名额', notes)

    def test_fingerprint_restores_single_mode_including_fixed_anchor(self):
        run = self.runs['combined']
        for row in run['ranking']:
            fingerprint = config_fingerprint(row)
            matches = find_fingerprint_matches(fingerprint, run['output'])
            self.assertEqual(len(matches), 1)
            restored = restore_fingerprint_match(matches[0])
            chosen = restored['selection']
            self.assertEqual(chosen['入场触发口径'], row['入场触发口径'])
            self.assertEqual(chosen['开仓条件']['1m'], [3])
            self.assertEqual(配置统计(chosen)['包含仓位完整组合数'], 1)
            self.assertEqual(config_fingerprint(matches[0]['row']), fingerprint)

    def test_chinese_front_and_detailed_definition_keep_original_machine_code(self):
        from openpyxl import load_workbook

        run = self.runs['combined']
        payload = json.loads((run['output'] / '各类止盈前5000名.json').read_text('utf-8'))
        before = copy.deepcopy(payload)
        view = build_view(payload, run['output'])
        self.assertEqual(payload, before)
        self.assertLess(view['表头'].index('入场次数规则'), view['表头'].index('入场触发口径'))
        native = {row['入场触发口径']: row for row in run['ranking']}
        category = next(key for key in view['分类'] if key.startswith('全局'))
        records = [dict(zip(view['表头'], row)) for row in view['分类'][category]]
        self.assertEqual({row['入场触发口径'] for row in records}, set(MODES))
        for row in records:
            mode = row['入场触发口径']
            self.assertEqual(row['入场次数规则'], mode_label(mode, 'entry'))
            self.assertIn('完整一轮', row['入场次数规则'])
            self.assertIn('初始截断色段不开仓', row['开仓判据（当前代码解释）'])
            self.assertIn('多空共用一次实际开仓名额', row['开仓判据（当前代码解释）'])
            self.assertIn('平仓均不恢复次数', row['开仓判据（当前代码解释）'])
            self.assertEqual(json.loads(row['完整单策略配置JSON'])['selection']['入场触发口径'], mode)
            self.assertEqual(row['策略指纹'], config_fingerprint(native[mode]))
        workbook_path = self.root / 'full_cycle.xlsx'
        export_portable_view(view, workbook_path)
        book = load_workbook(workbook_path, read_only=True)
        try:
            sheet = book[category]
            headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            actual = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
            self.assertEqual({row['入场触发口径'] for row in actual}, set(MODES))
            for row in actual:
                self.assertEqual(row['入场次数规则'], mode_label(row['入场触发口径'], 'entry'))
                self.assertEqual(json.loads(row['完整单策略配置JSON'])['selection']['入场触发口径'],
                                 row['入场触发口径'])
        finally:
            book.close()

    def test_legacy_fingerprint_algorithm_and_label_remain_unchanged(self):
        self.assertEqual(config_fingerprint({'入场触发口径': 'MACD_CYCLE'}), '61e306f2d621d852')
        self.assertEqual(mode_label('MACD_CYCLE', 'entry'), 'DIF/DEA金叉死叉每次红绿换色后最多一次')


if __name__ == '__main__':
    unittest.main()
