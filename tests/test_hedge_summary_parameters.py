"""Exported account/exit parameters must describe the actual normalized experiment."""
import copy
import csv
from datetime import datetime, timezone
from pathlib import Path
import tempfile
import unittest

from hedge_config import normalize_config
from hedge_signals import plan_entries
from hedge_worker import (CASE_HEADERS, SUMMARY_HEADERS, PARAMETER_HEADERS, CAPACITY_DATE_HEADER,
                          case_identity, summarize, summary_headers, csv_rows, export_excel)


class SummaryParameterTests(unittest.TestCase):
    def setUp(self):
        self.descriptor = plan_entries({}, {}, False)[0][0]
        self.stats = dict(ending_equity=23000., max_drawdown=.125, profit_factor=1.4,
                          trades=10., win_rate=.6, daily_trades=.5, fees=12.3, unrealized_pnl=-30.)

    def test_export_records_actual_normalized_parameters_and_keeps_identity(self):
        for mode, mode_label in [("single", "1倍单仓（不补仓）"), ("scale_in", "0.5倍首仓＋0.5倍补仓")]:
            with self.subTest(mode=mode):
                config = normalize_config(dict(mode=mode, initial_equity=23456.78, max_eth=17.5,
                                               maker_fee_rate=.000252, tp_weekday=.0125, tp_holiday=.0075))
                before = copy.deepcopy(config)
                identity = case_identity(self.descriptor, config, 12., .001, 1)
                summary = summarize(self.descriptor, config, 12., .001, 1, [self.stats])
                self.assertEqual(config, before)
                self.assertEqual(summary["策略标识"], identity)
                self.assertEqual(summary["本金（USDC）"], 23456.78)
                self.assertEqual(summary["持仓模式"], mode_label)
                self.assertEqual(summary["每方向名义倍数（倍）"], 1.)
                self.assertEqual(summary["每方向上限（ETH）"], 17.5)
                self.assertAlmostEqual(summary["Maker单边手续费（%）"], .0252)
                self.assertEqual(summary["工作日止盈（%）"], 1.25)
                self.assertEqual(summary["周末及中国节假日止盈（%）"], .75)
                self.assertEqual(summary["方向限制"], "未使用")
                self.assertEqual(summary["超时平仓（小时）"], 12.)

    def test_new_csv_appends_fields_and_accepts_legacy_summary_rows(self):
        config = normalize_config()
        descriptor = dict(self.descriptor, cases=(0, 3, 0, 57, 42))
        summary = summarize(descriptor, config, 6., .001, 0, [self.stats])
        old_headers = SUMMARY_HEADERS + CASE_HEADERS
        self.assertEqual(summary_headers([summary])[:len(old_headers)], old_headers)
        legacy = {name: value for name, value in summary.items() if name not in PARAMETER_HEADERS}
        self.assertEqual(summary_headers([legacy]), old_headers)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "summary.csv"
            csv_rows(path, summary_headers([summary]), [summary])
            with path.open(encoding="utf-8-sig", newline="") as stream:
                row = next(csv.DictReader(stream))
            self.assertEqual(row["策略标识"], summary["策略标识"])
            self.assertEqual(float(row["超时平仓（小时）"]), 6.)
            self.assertEqual(float(row["Maker单边手续费（%）"]), 0.)

    def test_excel_uses_fixed_seed_capacity_day_and_precise_fee_display(self):
        from openpyxl import load_workbook
        config = normalize_config(dict(mode="single", maker_fee_rate=.000252,
                                       direction_rule="yijing-midpoint-v1"))
        stats = dict(self.stats, direction_exits=2.)
        summary = summarize(self.descriptor, config, 6., 0., 1, [stats])
        original_summary = copy.deepcopy(summary)
        stamp = datetime(2025, 2, 12, 16, 1, tzinfo=timezone.utc).timestamp() * 1000
        fixed = dict(策略标识=summary["策略标识"], 开仓策略=summary["开仓策略"], 种子=7,
                     first_capacity_ms=stamp, **stats)
        notes = {"说明": "test"}
        with tempfile.TemporaryDirectory() as directory:
            target = export_excel(Path(directory), [summary], [fixed], [], notes)
            book = load_workbook(target, read_only=False, data_only=True)
            page = book["方案汇总"]
            headers = [cell.value for cell in page[1]]
            values = {name: page.cell(2, index+1).value for index, name in enumerate(headers)}
            self.assertEqual(headers[:10], ["策略指纹", "期末资金中位数（USDC）", "最大回撤中位数（%）",
                "每方向名义倍数（倍）", "胜率中位数（%）", "日均完整交易中位数", "成本模式",
                CAPACITY_DATE_HEADER, "每方向上限（ETH）", "利润因子PF中位数"])
            self.assertEqual(values["超时Maker退出（小时）"], 6.)
            self.assertNotIn("超时平仓（小时）", headers)
            self.assertEqual(values[CAPACITY_DATE_HEADER].date().isoformat(), "2025-02-13")
            self.assertEqual(values["成本模式"], "Maker 0.0252%/边")
            self.assertEqual(values["本金（USDC）"], 20000.)
            self.assertEqual(values["方向限制"], "红绿带中点")
            self.assertAlmostEqual(values["Maker单边手续费（%）"], .0252)
            self.assertEqual(page.cell(2, headers.index("Maker单边手续费（%）")+1).number_format, "0.0000")
            self.assertEqual(page.freeze_panes, "C2")
            explanations = dict(book["回测说明"].values)
            self.assertIn("未成交继续持仓", explanations["超时Maker退出"])
            self.assertIn("固定起始种子", explanations["首页首次达上限日期"])
            book.close()
        self.assertEqual(summary, original_summary)
        self.assertEqual(notes, {"说明": "test"})

    def test_missing_capacity_stat_is_not_reported_as_never_reached(self):
        from openpyxl import load_workbook
        summary = summarize(self.descriptor, normalize_config(), 6., 0., 0, [self.stats])
        with tempfile.TemporaryDirectory() as directory:
            target = export_excel(Path(directory), [summary], [], [], {})
            book = load_workbook(target, read_only=True, data_only=True)
            headers = next(book["方案汇总"].values)
            self.assertNotIn(CAPACITY_DATE_HEADER, headers)
            book.close()


if __name__ == "__main__":
    unittest.main()
