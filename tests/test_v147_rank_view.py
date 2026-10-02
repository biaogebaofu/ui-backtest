"""Check production export fields and portable format, not market performance."""
import copy
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook
from ranking_view import build_view, config_fingerprint
from export_runtime import export_xlsx_atomic
from test_v147_export_execution import sample_row


class ExecutionRankViewTests(unittest.TestCase):
    def payload(self, mixed=False):
        row = sample_row()
        row.update({'开仓MACD口径': 'MACD柱', '1分钟条件': '情况一',
                    '期末资金（USDC）':225, '累计收益率（%）':1.25,
                    '最大回撤（%）':.2, '名义倍数（倍）':2})
        records = [row]
        if mixed:
            records.append(dict(row, **{'平仓后最小开仓间隔（分钟）':0,
                                      '开仓净手续费率（%）':.0004}))
        return {'表头':list(row),'分类':{'移动止盈':[list(r.values()) for r in records]},
                '初始资金':100}

    def test_constant_execution_settings_and_mixed_rows_are_preserved(self):
        view = build_view(self.payload())
        settings = dict(view['共同设置'])
        self.assertEqual(settings['平仓后最小开仓间隔（分钟）'],5)
        self.assertEqual(settings['开仓净手续费率（%）'],.000252)
        self.assertEqual(settings['成本模式'],'LEGACY_COMBINED')
        mixed = build_view(self.payload(True))
        self.assertIn('开仓净手续费率（%）',mixed['表头'])
        self.assertIn('平仓后最小开仓间隔（分钟）',mixed['表头'])
        self.assertEqual(len(mixed['分类']['全局前5000']),2)

    def test_execution_changes_fingerprint(self):
        row=sample_row()
        for key in ('平仓后最小开仓间隔（分钟）','开仓净手续费率（%）',
                    '平仓净手续费率（%）','手续费返佣比例（%）','成本模式'):
            changed=copy.deepcopy(row);changed[key]=0
            self.assertNotEqual(config_fingerprint(row),config_fingerprint(changed))

    def test_portable_production_export_keeps_percent_units(self):
        import json
        root=Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory(prefix='v147_portable_') as tmp:
            folder=Path(tmp);source=folder/'ranking.json'
            source.write_text(json.dumps(self.payload()),'utf-8')
            # Verify the existing no-optional-renderer production path.
            with mock.patch.dict('sys.modules',{'artifact_tool':None}):
                output=export_xlsx_atomic(None,root/'excel_export.mjs',source,folder/'out.xlsx',root)
            wb=load_workbook(output,read_only=True,data_only=False)
            try:
                rows={r[0].value:r for r in wb['回测设置'].iter_rows()}
                fee=rows['开仓净手续费率（%）'][1]
                self.assertEqual(fee.value,.000252)
                self.assertEqual(fee.number_format,'0.00000%')
                self.assertEqual(rows['平仓后最小开仓间隔（分钟）'][1].value,5)
                self.assertIn('全局前5000',wb.sheetnames)
            finally:
                wb.close()


if __name__=='__main__':
    unittest.main()
