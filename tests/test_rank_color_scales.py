"""Display-only export regression: shared palettes, bounded ranges, unchanged data."""
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest import mock

from openpyxl import Workbook, load_workbook
from artifact_rank_export import build_workbook
from portable_rank_export import export_portable_view
from rank_color_scales import COLOR_SCALES, add_artifact_scales, add_openpyxl_scales
from streaming_excel_export import export_streaming_xlsx


ROOT = Path(__file__).resolve().parents[1]
HEADERS = ['期末资金（USDC）', '平均日完整交易数（次/日）', '胜率（%）', '名义倍数（倍）']
ROWS = [[100, 1, .1, 1], [200, 5, .5, 5], [300, 9, .9, 9]]


class RankColorScaleTests(unittest.TestCase):
    def assert_scales(self, sheet, first_row):
        rules = {str(cf.sqref): sheet.conditional_formatting[cf]
                 for cf in sheet.conditional_formatting}
        self.assertEqual(set(rules), {f'{c}{first_row}:{c}{first_row + 2}' for c in 'ABCD'})
        for column, spec in zip('ABCD', COLOR_SCALES):
            scale_rules = rules[f'{column}{first_row}:{column}{first_row + 2}']
            self.assertEqual(len(scale_rules), 1)
            self.assertEqual(scale_rules[0].type, 'colorScale')
            scale = scale_rules[0].colorScale
            self.assertEqual([(t.type, t.val) for t in scale.cfvo],
                             [('min', None), ('percentile', 50), ('max', None)])
            self.assertEqual([c.rgb[-6:] for c in scale.color],
                             [c.lstrip('#') for c in spec['colors']])

    def test_full_short_and_candidate_headers_share_all_four_scales(self):
        variants = [HEADERS,
                    ['期末资金', '日均完整交易', '胜率', '名义倍数'],
                    ['期末资金（USDC）', '日均完整交易（次/日）', '胜率（%）', '名义倍数（倍）'],
                    ['目标杠杆期末资金（USDC）', '平均日完整交易数（单/日）', '胜率（%）', '统一目标杠杆（倍）']]
        for headers in variants:
            with self.subTest(headers=headers):
                sheet = Workbook().active
                add_openpyxl_scales(sheet, headers, 5, 3)
                self.assert_scales(sheet, 5)

    def test_portable_low_middle_high_values_formulas_and_formats_unchanged(self):
        headers = HEADERS + ['账户净利润（USDC）', '初始资金（USDC）']
        rows = [row + [None, 50] for row in ROWS]
        view = {'表头': headers, '分类': {'最优': rows, '空榜': []}, '共同设置': []}
        with tempfile.TemporaryDirectory() as temp:
            target = Path(temp) / 'portable.xlsx'
            export_portable_view(view, target)
            workbook = load_workbook(target, data_only=False)
            try:
                sheet = workbook['最优']
                self.assert_scales(sheet, 2)
                self.assertEqual([list(row)[:4] for row in sheet.iter_rows(min_row=2, values_only=True)], ROWS)
                self.assertEqual([sheet[f'E{i}'].value for i in range(2, 5)], ['=A2-F2', '=A3-F3', '=A4-F4'])
                self.assertEqual(sheet['A2'].number_format, '#,##0.00;[Red]-#,##0.00')
                self.assertEqual(sheet['C2'].number_format, '0.00%;[Red]-0.00%')
                self.assertEqual(sheet['A1'].fill.fgColor.rgb[-6:], '17365D')
                self.assertEqual(sheet.auto_filter.ref, 'A1:F4')
                for name in ('回测设置', '字段说明', '空榜'):
                    self.assertEqual(len(workbook[name].conditional_formatting), 0)
            finally:
                workbook.close()

    def test_streaming_data_starts_on_row_five(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / 'input.json'
            target = Path(temp) / 'stream.xlsx'
            source.write_text(json.dumps({'表头': HEADERS, '分类': {'最差': ROWS}, '排行': '最差'}), 'utf-8')
            export_streaming_xlsx(source, target)
            workbook = load_workbook(target)
            try:
                sheet = workbook['最差']
                self.assert_scales(sheet, 5)
                self.assertEqual(list(sheet.values)[3], tuple(HEADERS))
                self.assertEqual([list(row) for row in list(sheet.values)[4:]], ROWS)
                self.assertEqual(sheet.auto_filter.ref, 'A4:D7')
                self.assertEqual(len(workbook['0_说明与检查'].conditional_formatting), 0)
            finally:
                workbook.close()

    def test_artifact_backend_uses_same_spec_and_data_ranges(self):
        sheets = {}
        workbook = mock.MagicMock()
        def add_sheet(name):
            sheets[name] = mock.MagicMock()
            return sheets[name]
        workbook.worksheets.add.side_effect = add_sheet
        api = SimpleNamespace(Workbook=SimpleNamespace(create=lambda: workbook))
        with mock.patch.dict(sys.modules, {'artifact_tool': api}), mock.patch('artifact_rank_export._write_table'):
            build_workbook({'表头': HEADERS, '分类': {'最优': ROWS, '空榜': []}})
        sheet = sheets['最优']
        self.assertEqual(sheet.get_range_by_indexes.call_args_list,
                         [mock.call(1, index, 3, 1) for index in range(4)])
        self.assertEqual(sheet.get_range_by_indexes.return_value.conditional_formats.add_color_scale.call_args_list,
                         [mock.call({'colors': s['colors'], 'thresholds': s['thresholds']}) for s in COLOR_SCALES[:4]])
        sheets['空榜'].get_range_by_indexes.assert_not_called()

    def test_empty_or_unrelated_columns_have_no_rules(self):
        sheet = Workbook().active
        add_openpyxl_scales(sheet, HEADERS, 2, 0)
        add_openpyxl_scales(sheet, ['盈亏比（倍）', '时间止损（分钟）'], 2, 3)
        self.assertEqual(len(sheet.conditional_formatting), 0)
        artifact_sheet = mock.MagicMock()
        add_artifact_scales(artifact_sheet, HEADERS, 2, 0)
        artifact_sheet.get_range_by_indexes.assert_not_called()

    @unittest.skipUnless(shutil.which('node'), 'Node is not installed')
    def test_javascript_adapter_and_exporter_wiring(self):
        script = '''
          const {addRankColorScales} = await import(process.argv[1]);
          const calls = [];
          const sheet = {getRangeByIndexes(...range) {return {conditionalFormats: {
            add(type, config) {calls.push({range, type, config});}
          }}}};
          addRankColorScales(sheet, JSON.parse(process.argv[2]), 5, 3);
          addRankColorScales(sheet, JSON.parse(process.argv[2]), 5, 0);
          console.log(JSON.stringify(calls));
        '''
        result = subprocess.run([shutil.which('node'), '--input-type=module', '-e', script,
                                 (ROOT / 'rank_color_scales.mjs').as_uri(), json.dumps(HEADERS)],
                                text=True, capture_output=True, check=True)
        calls = json.loads(result.stdout)
        self.assertEqual(calls, [{'range': [4, i, 3, 1], 'type': 'colorScale',
                                 'config': {'colors': s['colors'], 'thresholds': s['thresholds']}}
                                for i, s in enumerate(COLOR_SCALES[:4])])
        for exporter in ('excel_export.mjs', 'candidate_excel_export.mjs'):
            source = (ROOT / exporter).read_text('utf-8')
            self.assertIn('import { addRankColorScales } from "./rank_color_scales.mjs";', source)
            self.assertIn('addRankColorScales(sheet, headers, 5, rows.length);', source)


if __name__ == '__main__':
    unittest.main()
