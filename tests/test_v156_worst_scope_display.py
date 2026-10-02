"""Worst scope is explicit in both exporters, never inferred for old records."""
import copy
from pathlib import Path
import tempfile
import unittest
from unittest import mock

from ranking_view import build_view, ranking_scope_notes, config_fingerprint
from portable_rank_export import export_portable_view
import artifact_rank_export
import test_v147_rank_view as fixture


class WorstScopeDisplayTests(unittest.TestCase):
    def payload(self):
        payload = fixture.ExecutionRankViewTests().payload()
        payload.update(排行='最差', worst_scope='ALL_VALID_COMPLETED_V1',
                       最差排行范围='全部数值有效的已完成账户结果，不应用最优优选门槛；包含资金停机和归零账户',
                       组合门槛=[{'指标':'最大回撤（%）','条件':'<=','值':.2}])
        return payload

    def test_declared_scope_does_not_change_raw_identity_or_data(self):
        payload = self.payload(); before = copy.deepcopy(payload)
        view = build_view(payload)
        self.assertEqual(view['排行范围说明'],ranking_scope_notes(payload))
        notes = {row[0]:row[1] for row in view['排行范围说明']}
        self.assertIn('包含资金停机和归零账户',notes['最差排行范围'])
        self.assertIn('仅约束最优榜',notes['最优门槛适用范围'])
        self.assertEqual(notes['最优组合门槛'],payload['组合门槛'])
        raw = dict(zip(payload['表头'],payload['分类']['移动止盈'][0]))
        displayed = dict(zip(view['表头'],view['分类']['移动止盈'][0]))
        self.assertEqual(displayed['策略指纹'],config_fingerprint(raw))
        self.assertEqual(displayed['期末资金（USDC）'],raw['期末资金（USDC）'])
        self.assertEqual(payload,before)

    def test_missing_or_unrecognized_marker_cannot_claim_full_worst_scope(self):
        for marker in (None,'unknown'):
            payload = {'排行':'最差','最差排行范围':'不得采信的全量文字声明'}
            if marker:payload['worst_scope'] = marker
            notes = {row[0]:row[1] for row in ranking_scope_notes(payload)}
            self.assertEqual(notes['最差排行范围'],'旧记录未声明全量范围，建议从完整CSV重建')

    def test_both_exporters_use_same_prominent_scope_rows(self):
        from openpyxl import load_workbook
        view = build_view(self.payload())
        expected = view['排行范围说明']
        with tempfile.TemporaryDirectory(prefix='v156_scope_') as folder:
            output = Path(folder)/'scope.xlsx'
            export_portable_view(view,output)
            book = load_workbook(output,read_only=True)
            try:
                actual = list(book['回测设置'].iter_rows(min_row=2,max_row=1+len(expected),values_only=True))
                self.assertEqual(actual, [tuple(artifact_rank_export.clean(value) for value in row) for row in expected])
            finally:
                book.close()
        class Captured(Exception):pass
        tables = []
        def capture(sheet,matrix,name,**kwargs):
            tables.append(matrix)
            raise Captured()
        with mock.patch.dict('sys.modules',{'artifact_tool':mock.MagicMock()}), mock.patch.object(artifact_rank_export,'_write_table',side_effect=capture):
            with self.assertRaises(Captured):next(artifact_rank_export.workbook_steps(view))
        self.assertEqual(tables[0][1:1+len(expected)],expected)


if __name__ == '__main__':
    unittest.main()
