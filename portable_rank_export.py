"""Portable backend for the existing Windows UI (uses its installed openpyxl).

Called only when the optional artifact_tool renderer is unavailable. The original
CSV and raw ranking JSON stay untouched. See artifact_rank_export for the
artifact_tool implementation used to produce the delivered workbook.
"""
from __future__ import annotations
from pathlib import Path
from rank_color_scales import add_openpyxl_scales, COST_COLORS
from strategy_description import validate_cell_lengths
from account_statistics import CAPACITY_DATE


def export_portable_view(view,output_path):
    validate_cell_lengths(view)
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    from artifact_rank_export import clean,safe_name
    wb=Workbook(write_only=True);used={'回测设置','字段说明'}
    categories = view.get('分类', {})
    category_sheets = {name: wb.create_sheet(safe_name(name, used))
                       for name in sorted(categories, key=lambda name: not name.startswith('全局'))}
    header_font=Font(bold=True,color='FFFFFF',size=10);header_fill=PatternFill('solid',fgColor='17365D')
    alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    settings=wb.create_sheet('回测设置')
    header_cell=WriteOnlyCell(settings);header_cell.alignment=alignment
    header_cell.font=header_font;header_cell.fill=header_fill
    header_style=header_cell._style
    body_styles={}
    def body_style(number_format='General', fill=None):
        key=(number_format,fill)
        if key not in body_styles:
            cell=WriteOnlyCell(settings);cell.alignment=alignment;cell.number_format=number_format
            if fill:cell.fill=PatternFill('solid',fgColor=fill.lstrip('#'))
            body_styles[key]=cell._style
        return body_styles[key]
    def column_format(header):
        if header == CAPACITY_DATE:return 'yyyy-mm-dd'
        if header in ('策略指纹','基础策略编号') or 'JSON' in header:return '@'
        if any(word in header for word in ('费率','偏移','返佣')):return '0.000000%;[Red]-0.000000%'
        if '（%）' in header:return '0.00%;[Red]-0.00%'
        if 'USDC' in header:return '#,##0.00;[Red]-#,##0.00'
        if header in ('平均持仓（分钟）','利润因子PF','日均完整交易（次/日）'):return '0.00'
        return 'General'
    notes=[['项目','数值或说明','来源']]
    notes.extend(view.get('排行范围说明',[]))
    for k,v in view.get('本次扫描说明',{}).items():notes.append([k,clean(v),'本次运行记录'])
    for k,v in view.get('数据说明',{}).items():
        if k in ('bars','start_utc','end_utc'):notes.append([k,v,'回测数据说明.json'])
    constants={}
    for k,v in view.get('共同设置',[]):notes.append([k,v,'与主表一致，另集中列示']);constants[k]=len(notes)
    notes.append(['原始结果','全部回测结果.csv及原始排行JSON不改变','这里只调整阅读格式，不改变回测值'])
    def write(sheet,matrix,widths=None):
        sheet.freeze_panes='B2' if matrix[0][0]=='策略指纹' else 'A2'
        sheet.sheet_view.showGridLines=False
        sheet.auto_filter.ref=f'A1:{get_column_letter(len(matrix[0]))}{len(matrix)}'
        styles=[body_style(column_format(header)) for header in matrix[0]]
        for j,h in enumerate(matrix[0],1):
            width=(44 if h=='入场次数规则' else 18 if h=='成本模式' else 60 if 'JSON' in h else 55 if '判据' in h or h in ('成交与成本说明','定义来源与版本核对','参数完整性')
                   else 38 if h in ('开仓规则','信号/ATR止损','止盈规则') else 24 if h=='基础策略编号' else 20)
            width={'策略指纹':22,'期末资金（USDC）':20,'最大回撤（%）':14,'名义倍数（倍）':14,
                   '胜率（%）':13,'日均完整交易（次/日）':19,'成本模式':16,CAPACITY_DATE:32,
                   '开仓数量上限（ETH）':18,'排名':9}.get(h,width)
            sheet.column_dimensions[get_column_letter(j)].width=widths[j-1] if widths else width
        for i,row in enumerate(matrix,1):
            sheet.row_dimensions[i].height=40
            cells=[]
            for j,val in enumerate(row):
                if i>1 and matrix[0][j]=='基础策略编号' and val is not None:val=str(val)
                c=WriteOnlyCell(sheet,value=clean(val))
                c._style=header_style if i==1 else styles[j]
                if i>1 and matrix[0][j]=='成本模式' and val in COST_COLORS:
                    c._style=body_style(column_format(matrix[0][j]),COST_COLORS[val])
                if i>1 and sheet is settings and j==1 and '（%）' in str(row[0]):
                    c._style=body_style('0.00000%')
                cells.append(c)
            sheet.append(cells)
    write(settings,notes,[30,60,45])
    headers=view['表头']
    for category,rows in view.get('分类',{}).items():
        sheet=category_sheets[category];matrix=[headers]
        for i,row in enumerate(rows,2):
            row=list(row)
            if '账户净利润（USDC）' in headers:
                f=get_column_letter(headers.index('期末资金（USDC）')+1)
                if '初始资金（USDC）' in headers:initial=f'{get_column_letter(headers.index("初始资金（USDC）")+1)}{i}'
                elif '初始资金（USDC）' in constants:initial=f"'回测设置'!$B${constants['初始资金（USDC）']}"
                else:raise ValueError('缺少初始资金，不能生成净利润公式')
                row[headers.index('账户净利润（USDC）')]=f'={f}{i}-{initial}'
            matrix.append(row)
        write(sheet,matrix)
        add_openpyxl_scales(sheet,headers,2,len(rows))
    dic=wb.create_sheet('字段说明');write(dic,[['字段','解释']]+view.get('字段说明',[]),[34,80])
    if view.get('止盈方案字典'):
        tpdict=wb.create_sheet('止盈方案字典')
        write(tpdict,[['原编号','类别','周期','指标','参数一','参数二','参数三','完整说明及单位']]+view['止盈方案字典'],[10,24,12,22,16,16,16,72])
    output_path=Path(output_path);output_path.parent.mkdir(parents=True,exist_ok=True);wb.save(output_path)
