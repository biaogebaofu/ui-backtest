"""Offline UI worker for random/individual-entry hedge strategy comparisons."""
from __future__ import annotations

import argparse
from collections import defaultdict
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
import csv
from datetime import datetime, timedelta, timezone
import hashlib
import itertools
import json
import math
import os
from pathlib import Path
import shutil
import sys
import threading
import time
import traceback

import numpy as np

BASE = Path(__file__).resolve().parent
BJ = timezone(timedelta(hours=8))
DAY = 86400000
SUMMARY_HEADERS = ["策略标识", "开仓策略", "期末资金中位数（USDC）", "最大回撤中位数（%）", "利润因子PF中位数",
                   "完整交易次数中位数", "胜率中位数（%）", "日均完整交易中位数", "超时平仓（小时）", "补仓触发跌幅（%）",
                   "K线路径", "随机重复次数", "盈利比例（%）", "期末资金P10", "期末资金P90", "最差最大回撤（%）",
                   "费用中位数（USDC）", "未平仓损益中位数", "开仓周期", "开仓代码", "MACD口径", "入场次数口径",
                   "方向限制", "方向切换平仓中位数"]
CASE_HEADERS = [f"{tf}条件代码" for tf in ("4h", "1h", "15m", "5m", "1m")]
PARAMETER_HEADERS = ["本金（USDC）", "持仓模式", "每方向名义倍数（倍）", "每方向上限（ETH）",
                     "Maker单边手续费（%）", "工作日止盈（%）", "周末及中国节假日止盈（%）", "退出成交方式", "成本模式"]
CAPACITY_DATE_HEADER = "首次达方向上限日期（首种子UTC+8）"
PF_DEFINITION = ("双向PF＝完整交易的正净利润金额合计÷负净利润金额绝对值合计；计入该交易首仓、补仓、减仓及平仓手续费。"
                 "期末未平仓损益不计入PF。普通回测旧PF采用逐笔净收益率正和÷负和，口径不同，不宜直接横比。")


class Stopped(Exception):
    pass


def clean(value):
    if isinstance(value, dict):
        return {key: clean(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean(item) for item in value]
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def emit(kind, **payload):
    print(json.dumps(clean(dict(type=kind, **payload)), ensure_ascii=False, allow_nan=False), flush=True)


def write_json(path, payload):
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(clean(payload), ensure_ascii=False, indent=2, allow_nan=False), "utf-8")
    os.replace(temporary, path)


def dynamic_tp(timestamps, config, calendar_path=None):
    calendar_path = Path(calendar_path or BASE / "hedge_holidays.json")
    calendar = json.loads(calendar_path.read_text("utf-8"))
    days = (timestamps + 8 * 3600000) // DAY
    years = np.unique(days.astype("datetime64[D]").astype("datetime64[Y]").astype(np.int64) + 1970)
    missing = [str(year) for year in years if str(year) not in calendar["years"]]
    if missing:
        raise ValueError("缺少中国法定节假日日历年份：" + "、".join(missing) + "；请先补充已核实日历，不能静默当作普通工作日")
    holiday_days = set()
    for year in years:
        for span in calendar["years"][str(year)]:
            first = int(np.datetime64(span["start"], "D").astype(np.int64))
            last = int(np.datetime64(span["end"], "D").astype(np.int64))
            holiday_days.update(range(first, last + 1))
    holidays = ((days + 3) % 7 >= 5) | np.isin(days, np.asarray(sorted(holiday_days), dtype=np.int64))
    return np.where(holidays, config["tp_holiday"], config["tp_weekday"]).astype(np.float64)


def load_data(feature_path, config):
    with np.load(feature_path, allow_pickle=False) as features:
        data = {out: np.asarray(features[name]) for out, name in
                (("ts", "openms"), ("o", "open"), ("h", "high"), ("l", "low"), ("c", "close"))}
    ts = data["ts"]
    if len(ts) < 2 or not np.issubdtype(ts.dtype, np.integer) or np.any(np.diff(ts) != 60000) or np.any(ts % 60000):
        raise ValueError("双向回测要求时间戳为UTC毫秒、连续完整的1分钟K线；发现缺失、重复或不正确单位，未补造K线")
    if any(len(values) != len(ts) or np.any(~np.isfinite(values)) or np.any(values <= 0)
           for name, values in data.items() if name != "ts"):
        raise ValueError("K线OHLC存在缺失、非有限值、非正价格或长度不一致")
    if np.any(data["h"] < np.maximum(data["o"], data["c"])) or np.any(data["l"] > np.minimum(data["o"], data["c"])):
        raise ValueError("K线OHLC高低价关系不成立")
    data["tp"] = dynamic_tp(ts, config)
    if config.get("direction_rule"):
        from direction_calendar import direction_array
        data["direction_limits"] = direction_array(np.r_[ts, ts[-1]+60000])
    for values in data.values():
        values.flags.writeable = False
    return data


def case_identity(descriptor, config, hours, drop, path):
    from hedge_engine import ENGINE_VERSION
    from hedge_signals import strategy_id
    account = {key: value for key, value in config.items()
               if key not in ("seeds", "seed_start", "compare_entries", "add_drops", "timeout_hours", "paths")}
    return strategy_id(dict(engine_version=ENGINE_VERSION, entry=descriptor, config=account,
                            hours=hours, add_drop=drop, path=path))


def summarize(descriptor, config, hours, drop, path, rows):
    from hedge_config import normalize_config
    account = normalize_config(config)
    def quantile(key, q=.5):
        values = np.array([row.get(key, np.nan) for row in rows], dtype=float)
        # Positive infinity is meaningful for PF without losses; do not turn it into zero.
        values = values[~np.isnan(values)]
        if not len(values):
            return None
        return float(np.median(values) if q == .5 else np.quantile(values, q))
    identity = case_identity(descriptor, config, hours, drop, path)
    values = [identity, descriptor["label"], quantile("ending_equity"), quantile("max_drawdown") * 100,
              quantile("profit_factor"), quantile("trades"),
              None if quantile("win_rate") is None else quantile("win_rate") * 100,
              quantile("daily_trades"), hours, drop * 100, "开高低收" if path == 0 else "开低高收", len(rows),
              sum(row["ending_equity"] > config["initial_equity"] for row in rows) / len(rows) * 100,
              quantile("ending_equity", .1), quantile("ending_equity", .9), quantile("max_drawdown", 1) * 100,
              quantile("fees"), quantile("unrealized_pnl"), descriptor["timeframe"], descriptor["code"],
              descriptor["field"], descriptor["entry_mode"],
              "红绿带中点" if config.get("direction_rule") else "未使用",
              quantile("direction_exits") if config.get("direction_rule") else 0]
    row = dict(zip(SUMMARY_HEADERS, values))
    if sum(bool(code) for code in descriptor["cases"]) > 1:
        row.update(zip(CASE_HEADERS, descriptor["cases"]))
    row.update(zip(PARAMETER_HEADERS, [account["initial_equity"],
                   "1倍单仓（不补仓）" if account["mode"] == "single" else "0.5倍首仓＋0.5倍补仓",
                   account["first_multiple"] + account["add_multiple"], account["max_eth"],
                   account["maker_fee_rate"] * 100, account["tp_weekday"] * 100,
                   account["tp_holiday"] * 100, "Maker挂单，未来严格穿价成交",
                   f"Maker {account['maker_fee_rate'] * 100:g}%/边"]))
    return row


def summary_headers(rows):
    # Existing names and positions stay intact for old CSV readers/fingerprint imports.
    return (SUMMARY_HEADERS + (CASE_HEADERS if any(CASE_HEADERS[0] in row for row in rows) else [])
            + [name for name in PARAMETER_HEADERS if any(name in row for row in rows)])


def seed_independent(signals, config):
    # The engine reads coins[i] only when both flat sides are permitted together.
    # A unique calendar direction or nonoverlapping permissions makes that unreachable.
    return bool(config.get("direction_rule")) or not np.any(signals[:, 0] & signals[:, 1])


def csv_rows(path, headers, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def export_excel(output, summaries, fixed_rows, skipped, notes):
    """Small, readable overview; complete repetitions remain in the streamed CSV."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook(write_only=True)
    def sheet(title, headers, rows):
        page = workbook.create_sheet(title)
        page.freeze_panes = "C2"
        page.sheet_view.showGridLines = False
        for index, header in enumerate(headers, 1):
            page.column_dimensions[get_column_letter(index)].width = 88 if header == "开仓策略" else (22 if index == 1 else 19)
        head = []
        for index, header in enumerate(headers):
            label = {"超时平仓（小时）": "超时Maker退出（小时）", "策略标识": "策略指纹"}.get(header, header)
            cell = WriteOnlyCell(page, value=label)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="156572" if header in
                ("期末资金中位数（USDC）", "最大回撤中位数（%）", "利润因子PF中位数", "完整交易次数中位数", "胜率中位数（%）", "日均完整交易中位数") else "17324D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            head.append(cell)
        page.append(head)
        page.row_dimensions[1].height = 48
        count = 1
        for row in rows:
            cells = []
            for index, header in enumerate(headers):
                value = row.get(header)
                # Rule codes are identifiers; Excel numbers keep only 15 significant digits.
                if header == "开仓代码" or header in CASE_HEADERS:
                    value = None if value is None else str(value)
                if isinstance(value, (float, np.floating)) and not math.isfinite(value):
                    value = "∞" if value > 0 else None
                cell = WriteOnlyCell(page, value=value)
                if header == "开仓策略":
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                    width = sum(2 if ord(char) > 255 else 1 for char in str(value or ""))
                    page.row_dimensions[count + 1].height = max(20, 16 * math.ceil(width / 82))
                elif header in ("退出成交方式", "持仓模式"):
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                    page.row_dimensions[count + 1].height = max(page.row_dimensions[count + 1].height or 20, 32)
                if isinstance(value, (float, np.floating)):
                    cell.number_format = "#,##0.00"
                if header in ("Maker单边手续费（%）", "工作日止盈（%）", "周末及中国节假日止盈（%）"):
                    cell.number_format = "0.0000"
                elif header == CAPACITY_DATE_HEADER and not isinstance(value, str):
                    cell.number_format = "yyyy-mm-dd"
                if header == "期末资金中位数（USDC）":
                    cell.fill = PatternFill("solid", fgColor="E5F3E9")
                elif "回撤" in header:
                    cell.fill = PatternFill("solid", fgColor="FFF0DA")
                elif "PF" in header:
                    cell.fill = PatternFill("solid", fgColor="E5EFF9")
                cells.append(cell)
            page.append(cells)
            count += 1
        page.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count}"
    summary_view = [dict(row) for row in summaries]
    fixed_by_id = {row["策略标识"]: row for row in fixed_rows}
    for row in summary_view:
        stamp = fixed_by_id.get(row["策略标识"], {}).get("first_capacity_ms")
        if stamp is not None:
            row[CAPACITY_DATE_HEADER] = datetime.fromtimestamp(stamp / 1000, BJ).date() if stamp >= 0 else "未达到"
    headers = summary_headers(summaries)
    if any(CAPACITY_DATE_HEADER in row for row in summary_view):
        headers.append(CAPACITY_DATE_HEADER)
    front = ["策略标识", "期末资金中位数（USDC）", "最大回撤中位数（%）", "每方向名义倍数（倍）",
             "胜率中位数（%）", "日均完整交易中位数", "成本模式", CAPACITY_DATE_HEADER,
             "每方向上限（ETH）", "利润因子PF中位数", "完整交易次数中位数", "本金（USDC）",
             "持仓模式", "工作日止盈（%）", "周末及中国节假日止盈（%）", "方向限制", "超时平仓（小时）"]
    headers = [name for name in front if name in headers] + [name for name in headers if name not in front]
    sheet("方案汇总", headers, summary_view)
    if fixed_rows:
        names = {"ending_equity": "期末资金（USDC）", "net_profit": "净利润（USDC）", "max_drawdown": "最大回撤（%）",
                 "超时小时": "超时Maker退出（小时）",
                 "trades": "完整交易次数", "wins": "盈利次数", "win_rate": "胜率（%）", "profit_factor": "利润因子PF",
                 "daily_trades": "日均完整交易", "fees": "手续费（USDC）", "time_exits": "超时平仓次数", "tp_exits": "止盈次数",
                 "direction_exits": "方向切换平仓次数",
                 "direction_pending_eth": "期末待方向退出数量（ETH）",
                 "average_hold_minutes": "平均持仓（分钟）", "max_hold_minutes": "最长持仓（分钟）",
                 "max_timeout_delay_minutes": "超时成交最大延迟（分钟）", "first_capacity_ms": "首次达到方向数量上限（北京时间）",
                 "open_positions": "期末持仓方向数", "unrealized_pnl": "期末未实现损益", "min_equity": "最低权益",
                 "entries": "首仓及补仓总次数", "min_entry_gap_minutes": "实际最短开仓间隔（分钟）",
                 "capacity_entries": "达到方向数量上限次数", "final_cash": "期末账面余额", "add_entries": "补仓次数",
                 "breakeven_exits": "保本减仓次数", "return_rate": "累计收益率（%）", "max_direction_eth": "最大单方向数量（ETH）",
                 "ending_long_eth": "期末多仓数量（ETH）", "ending_short_eth": "期末空仓数量（ETH）"}
        fixed_view = []
        for row in fixed_rows:
            values = {}
            for key, value in row.items():
                if key in ("max_drawdown", "win_rate", "return_rate"):
                    value *= 100
                elif key == "first_capacity_ms":
                    value = datetime.fromtimestamp(value/1000, BJ).isoformat() if value >= 0 else "未达到"
                elif key == "路径":
                    value = "开高低收" if value == 0 else "开低高收"
                elif key == "补仓触发比例":
                    key, value = "补仓触发跌幅（%）", value * 100
                values[names.get(key, key)] = value
            fixed_view.append(values)
        heads = ["策略标识", "开仓策略", "期末资金（USDC）", "最大回撤（%）", "利润因子PF", "完整交易次数", "胜率（%）"]
        heads += [key for key in fixed_view[0] if key not in heads]
        sheet("固定种子汇总", heads, fixed_view)
    sheet("自动跳过清单", ["周期", "开仓代码", "开仓规则", "跳过原因"], skipped)
    notes = dict(notes)
    notes["利润因子PF口径"] = PF_DEFINITION
    notes["超时Maker退出"] = "从首次入场起计时；到期提交Maker平仓挂单，之后严格穿价才成交，未成交继续持仓。"
    if any(CAPACITY_DATE_HEADER in row for row in summary_view):
        notes["首页首次达上限日期"] = "固定起始种子的首次实际达到每方向数量上限日期（UTC+8）；不是各随机种子日期的中位数。"
    sheet("回测说明", ["项目", "说明"], [{"项目": str(key), "说明": str(value)} for key, value in notes.items()])
    target = output / ("双向策略比较_" + datetime.now(BJ).strftime("%Y%m%d_%H%M%S") + ".xlsx")
    temporary = target.with_name("." + target.stem + ".tmp.xlsx")
    try:
        workbook.save(temporary)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return target


def execute(request, output, threads):
    from feature_builder import build_features, 特征版本
    from data_sources import source_bundle_fingerprint
    from hedge_config import normalize_config
    from hedge_engine import run_case, ENGINE_VERSION as HEDGE_VERSION
    from hedge_signals import SignalAdapter
    from hedge_entry_combinations import plan_request_entries, fifth_tasks
    from run_safety import build_run_identity
    from account_statistics import ENGINE_VERSION
    from indicator_combinations import DEFAULT_REGISTRY

    config = normalize_config(request.get("hedge", {}))
    selection = request.get("selection", {})
    if request.get("indicator_registry"):
        DEFAULT_REGISTRY.update(request["indicator_registry"])
    stop_file = output / "停止请求.flag"
    control = output / "控制"
    control.mkdir(exist_ok=True)
    bridge_done = threading.Event()
    def bridge():
        while not bridge_done.wait(.2):
            if stop_file.exists():
                (control / "停止.flag").touch()
                return
    def stop():
        if stop_file.exists():
            raise Stopped()
        while (control / "暂停.flag").exists():
            if stop_file.exists():
                raise Stopped()
            time.sleep(.1)
    monitor = threading.Thread(target=bridge, daemon=True)
    monitor.start()
    try:
        stop()
        sources = {name: request.get("sources", {}).get(name, "") for name in ("kline", "micro", "funding", "oi", "bundle")}
        if not sources["kline"] and not sources["bundle"]:
            raise ValueError("请先在数据源页选择1分钟K线或数据包")
        feature_request = dict(fingerprint=source_bundle_fingerprint(sources), start=request.get("start", ""),
                               end=request.get("end", ""), feature_version=特征版本)
        if request.get("hedge_exact") and feature_request != request["hedge_exact"]["feature_request"]:
            raise ValueError("行情文件或日期已变化，未按其他数据运行该双向指纹")
        cache_key = hashlib.sha256(json.dumps(feature_request, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:20]
        cache = Path(request.get("cache_root") or output.parent / "_共享指标缓存") / cache_key
        emit("stage", message="正在读取主界面选定行情及日期，检查共享指标缓存；双向回测使用独立账户模型")
        feature_path = build_features(sources["kline"], str(cache), request.get("start"), request.get("end"), False,
                                      sources["micro"], sources["funding"], sources["oi"], sources["bundle"])
        meta = json.loads(feature_path.with_name("features_meta.json").read_text("utf-8"))
        stop()
        data = load_data(feature_path, config)
        if request.get("hedge_exact"):
            from hedge_fingerprints import resolve_exact
            descriptor, variant = resolve_exact(request, meta)
            descriptors, skipped = [descriptor], []
        else:
            descriptors, skipped = plan_request_entries(dict(request, hedge=config), meta)
        # Persist any same-timeframe groups generated by the selection plan.
        request = dict(request, indicator_registry=DEFAULT_REGISTRY.to_dict())
        write_json(output / "双向回测请求.json", dict(request, hedge=config))
        write_json(output / "双向运行身份.json", dict(engine_version=HEDGE_VERSION, feature_request=meta["request"],
                   files_sha256={name: hashlib.sha256((BASE / name).read_bytes()).hexdigest() for name in
                                 ("hedge_engine.py", "hedge_config.py", "hedge_signals.py", "hedge_worker.py", "hedge_holidays.json",
                                  "hedge_entry_combinations.py") + (("direction_calendar.py",) if config.get("direction_rule") else ())}))
        write_json(output / "行情与节假日核对.json", dict(features=meta, rows=len(data["ts"]),
                   calendar_sha256=hashlib.sha256((BASE / "hedge_holidays.json").read_bytes()).hexdigest(),
                   first_open_ms=int(data["ts"][0]), end_exclusive_ms=int(data["ts"][-1] + 60000)))
        write_json(output / "开仓独立比较清单.json", dict(entries=descriptors, skipped=skipped))
        shutil.copyfile(BASE / "hedge_holidays.json", output / "已用节假日日历.json")
        if config.get("direction_rule"):
            from direction_calendar import event_window
            first=datetime.fromtimestamp(int(data['ts'][0])/1000,BJ)
            last=datetime.fromtimestamp(int(data['ts'][-1]+60000)/1000,BJ)
            events=event_window(first,last)
            calendar_rows=[]
            for index,event in enumerate(events):
                finish=events[index+1][0] if index+1<len(events) else last
                calendar_rows.append({"开始（含）":max(first,event[0]).isoformat(),"结束（不含）":finish.isoformat(),
                    "允许方向":"只多" if event[1]==1 else "只空","中点":event[2].isoformat(),
                    "原生效时间":event[0].isoformat()})
            write_json(output/"已用方向限制日历.json",dict(rule=config["direction_rule"],timezone="UTC+08:00",intervals=calendar_rows))
            shutil.copyfile(BASE/"direction_calendar.py",output/"已用方向限制公式.py")
            emit("stage",message=f"已使用红绿带中点方向限制：{len(events)}段；历史年份按公式重算；切换先Maker平旧仓，再允许新方向")
        else:
            emit("stage",message="未使用方向限制：沿用原双向持仓规则")
        for row in skipped:
            emit("stage", message=f"已跳过 {row['周期']} {row['开仓规则']}：{row['跳过原因']}")
        warmup = 200 if config["compare_entries"] else 0
        adapter = None
        rule_count = sum(row["id"] != "RANDOM" for row in descriptors)
        if rule_count:
            label = "按所选周期组合，所有启用条件共同过滤" if request.get("entry_plan") == "combined" else "单周期独立对照"
            emit("stage", message=f"准备{rule_count}条原开仓规则；{label}；共同预热200根")
            adapter = SignalAdapter(feature_path)
            tasks = fifth_tasks(descriptors)
            if tasks:
                from fifth_precompute import prepare_fifth_signals, PrecomputeStopped
                identity = build_run_identity(BASE, "hedge-independent", meta, ENGINE_VERSION)
                def precompute_progress(details):
                    completed = details["completed"]
                    total = details["total"]
                    eta = details.get("elapsed_seconds", 0) / completed * (total-completed) if completed else None
                    emit("progress", completed=completed, total=total, eta_seconds=eta,
                         message="[开仓指标预计算] " + details["message"])
                try:
                    cache_path = prepare_fifth_signals(adapter.engine.E.D, tasks, output, identity, threads,
                        precompute_progress, control, cache / "第五轮信号缓存")
                except PrecomputeStopped:
                    raise Stopped() from None
                adapter.engine.set_fifth_signal_cache(cache_path)
        drops = config["add_drops"] if config["mode"] == "scale_in" else [0.]
        variants = list(itertools.product(config["timeout_hours"], drops, config["paths"]))
        if request.get("hedge_exact"):
            variants = [variant]
        total = len(descriptors) * len(variants) * config["seeds"]
        completed = 0
        computed_cases = 0
        summaries, fixed_rows = [], []
        replay_started = time.monotonic()
        last_progress = 0.
        interrupted = False
        raw_path = output / "全部独立试验.csv"
        if raw_path.exists():
            raise ValueError("结果目录已有双向回测CSV，请使用新目录，避免覆盖或混写历史结果")
        emit("stage", message=f"开始{len(descriptors)}种开仓 × {len(variants)}组退出/补仓/路径 × {config['seeds']}次随机重复，共{total:,}次；Maker未来穿价成交")
        with raw_path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = None
            with ThreadPoolExecutor(max_workers=threads) as pool:
                for descriptor in descriptors:
                    if stop_file.exists():
                        interrupted = True
                        break
                    stop()
                    emit("stage", message="准备开仓许可：" + descriptor["label"])
                    if descriptor["id"] == "RANDOM":
                        signals = np.ones((len(data["ts"]), 2), dtype=np.bool_)
                        signals[:warmup] = False
                        signals.flags.writeable = False
                        ids = np.full(len(data["ts"]), -1, dtype=np.int64)
                        ids.flags.writeable = False
                    else:
                        signals, ids = adapter.build(descriptor, warmup)
                    grouped = defaultdict(list)
                    equivalent_seeds = seed_independent(signals, config)
                    seeds = range(config["seed_start"], config["seed_start"] + config["seeds"])
                    calculation_seeds = (config["seed_start"],) if equivalent_seeds else seeds
                    if equivalent_seeds and config["seeds"] > 1:
                        emit("stage", message=f"该开仓规则没有随机选向分支：每组参数实际计算1次，保留{config['seeds']}个等价种子记录；不是独立样本")
                    keys = iter(itertools.product(variants, calculation_seeds))
                    futures = {}
                    def submit_one():
                        try:
                            (hours, drop, path), seed = next(keys)
                        except StopIteration:
                            return False
                        while (control / "暂停.flag").exists() and not stop_file.exists():
                            time.sleep(.1)
                        if stop_file.exists():
                            return False
                        future = pool.submit(run_case, data, config, hours, seed, signals=signals, run_ids=ids,
                                             add_drop=drop, path=path, detail=False)
                        futures[future] = (hours, drop, path, seed)
                        return True
                    for _ in range(max(1, threads * 2)):
                        if not submit_one():
                            break
                    while futures:
                        done, _ = wait(futures, timeout=.5, return_when=FIRST_COMPLETED)
                        if stop_file.exists():
                            interrupted = True
                            for future in futures:
                                future.cancel()
                        for future in done:
                            hours, drop, path, seed = futures.pop(future)
                            if future.cancelled():
                                continue
                            result = future.result()
                            computed_cases += 1
                            stats = {key: float(value) for key, value in result["stats"].items()}
                            stats.update(ending_long_eth=float(result["positions"][0, 1]),
                                         ending_short_eth=float(result["positions"][1, 1]))
                            identity = case_identity(descriptor, config, hours, drop, path)
                            for saved_seed in seeds if equivalent_seeds else (seed,):
                                if stop_file.exists():
                                    interrupted = True
                                    break
                                row = dict(策略标识=identity, 开仓策略=descriptor["label"], 超时小时=hours,
                                           补仓触发比例=drop, 路径=path, 种子=saved_seed,
                                           方向限制="红绿带中点" if config.get("direction_rule") else "未使用", **stats)
                                if writer is None:
                                    writer = csv.DictWriter(stream, fieldnames=list(row))
                                    writer.writeheader()
                                writer.writerow(row)
                                grouped[(hours, drop, path)].append(stats)
                                if saved_seed == config["seed_start"]:
                                    fixed_rows.append(row)
                                completed += 1
                            if not interrupted:
                                submit_one()
                        now = time.monotonic()
                        if now-last_progress >= 1 or not futures:
                            elapsed = now-replay_started
                            eta = elapsed/completed*(total-completed) if completed else None
                            emit("progress", completed=completed, total=total, eta_seconds=eta,
                                 message=f"{descriptor['label']}；已完成 {completed:,}/{total:,} 次" + ("；正在安全停止" if interrupted else ""))
                            stream.flush()
                            last_progress = now
                    for (hours, drop, path), rows in grouped.items():
                        summaries.append(summarize(descriptor, config, hours, drop, path, rows))
                    if interrupted:
                        break
        summaries.sort(key=lambda row: row["期末资金中位数（USDC）"], reverse=True)
        csv_rows(output / "方案汇总.csv", summary_headers(summaries), summaries)
        csv_rows(output / "自动跳过清单.csv", ["周期", "开仓代码", "开仓规则", "跳过原因"], skipped)
        notes = {"状态": "安全停止，以下仅含已完成试验" if interrupted else "全部完成", "模型版本": HEDGE_VERSION,
                 "利润因子PF口径": PF_DEFINITION,
                 "方向限制": ("红绿带中点：中点对应2H收盘生效，历史年份重算；禁止反向首仓和补仓；切换撤销旧开仓单并Maker平旧仓，成交确认前禁止新方向；允许空仓"
                             if config.get("direction_rule") else "未使用，沿用原双向规则"),
                 "方向切换成交": "切换并不保证立即成交，沿用Maker未来严格穿价；方向切换退出单独计数，不混入时间止损；K线内成交时间沿用该分钟收盘时间记账",
                 "本金": config["initial_equity"], "双向模式": config["mode"], "每方向最大ETH": config["max_eth"],
                 "Maker手续费率": config["maker_fee_rate"], "资金费率": "未计；本轮不读取或推测历史资金费成本",
                 "保证金与强平": "未模拟真实交易所逐档维持保证金、强平手续费和执行；无固定比例止损不代表真实账户可忽略强平风险",
                 "成交假设": "仅Maker；报价之后K线路径严格穿价才模拟成交，不保证真实盘口排队成交，不使用吃单兜底",
                 "止盈挂单": "实际首仓/补仓后立即按当时已知数据设置预挂退出；仅允许之后的K线路径触发，同根入场前价格不参与",
                 "双K线路径": "开高低收与开低高收独立列出，比较分钟内先后次序不确定性",
                 "到期": "按首次入场计时；到期提交Maker平仓单，未穿价继续持仓，并记录延迟；样本结束持仓按市价计权益",
                 "入场间隔": f"任意方向实际首仓和补仓共用{config['entry_gap_minutes']}分钟间隔",
                 "开仓比较": "逐条选中规则；保留原MACD口径、入场次数、S3条件；方向BOTH、会话ALL、位置OFF；其他旧止盈止损不继承",
                 "组合范围": ("各周期选择逐一交叉组合；同周期的指标组合沿用第2页设置；所有启用周期共同过滤，触发与次数口径沿用原引擎"
                             if request.get("entry_plan") == "combined" else "单周期独立对照；各周期间不交叉相乘"),
                 "预热": f"所有方案前{warmup}根禁止首仓",
                 "随机重复": f"从种子{config['seed_start']}开始；各条件及参数共用按分钟索引的相同随机流；固定种子汇总不挑幸运种子",
                 "等价种子计算": f"实际运行{computed_cases}次计算，输出{completed}条种子记录；仅在方向日历开启或开仓许可从不同时允许多空时复用统计。等价记录不是独立市场样本。",
                 "Excel与CSV": "首页按期末资金中位数排序；所有重复试验在全部独立试验.csv，Excel另列固定首个种子汇总；这些均是试验级汇总，不是逐笔交易明细",
                 "策略标识": "双向模型指纹；可在第2页多指纹列表导入，绑定原持仓模型及随机种子范围独立重测",
                 "日历": "中国法定节假日+每周六日，北京时间；补班周末仍使用周末止盈比例，跨日动态调整",
                 "日期范围": f"{datetime.fromtimestamp(data['ts'][0]/1000, BJ).isoformat()} 至 {datetime.fromtimestamp((data['ts'][-1]+60000)/1000, BJ).isoformat()}",
                 "完整配置JSON": json.dumps(config, ensure_ascii=False)}
        write_json(output / "回测说明.json", notes)
        write_json(output / "回测完成状态.json", dict(completed=completed, total=total, stopped=interrupted))
        emit("stage", message=f"实际计算{computed_cases:,}次，保留{completed:,}条种子记录；CSV已保存，正在生成方案汇总Excel和固定种子对照")
        workbook = export_excel(output, summaries, fixed_rows, skipped, notes)
        emit("stopped" if interrupted else "done", output=str(output), workbook=str(workbook),
             completed=completed, total=total, message="已安全停止并导出已完成结果" if interrupted else "双向策略回测及导出已完成")
    finally:
        bridge_done.set()
        monitor.join(timeout=1)


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser()
    parser.add_argument("--request", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--threads", type=int, default=max(1, (os.cpu_count() or 2)-2))
    args = parser.parse_args()
    output = Path(args.output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    try:
        if not 1 <= args.threads <= max(1, os.cpu_count() or 1):
            raise ValueError("并行线程数超出当前CPU可用范围")
        from run_safety import exclusive_output
        with exclusive_output(output):
            execute(json.loads(Path(args.request).read_text("utf-8-sig")), output, args.threads)
    except Stopped:
        emit("stopped", output=str(output), message="已安全停止；此前完成的缓存和结果保留")
    except Exception as exc:
        (output / "错误详情.txt").write_text(traceback.format_exc(), "utf-8")
        emit("error", output=str(output), message=str(exc))
        raise SystemExit(1) from None


if __name__ == "__main__":
    main()
