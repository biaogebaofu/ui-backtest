"""Shared display-only scales; frequency and leverage do not imply quality."""
import json
from pathlib import Path


COLOR_SCALES = json.loads(Path(__file__).with_suffix('.json').read_text('utf-8'))
COST_COLORS = {'仅手续费': '#DDEBF7', 'FEE': '#DDEBF7',
               '仅成交偏移': '#FCE4D6', 'SLIPPAGE': '#FCE4D6'}


def matching_scales(headers):
    for index, header in enumerate(headers):
        for spec in COLOR_SCALES:
            if header in spec['headers']:
                yield index, spec
                break


def add_openpyxl_scales(sheet, headers, first_row, row_count):
    if row_count < 1:
        return
    from openpyxl.formatting.rule import ColorScale, FormatObject, Rule, CellIsRule
    from openpyxl.styles import Color, PatternFill
    from openpyxl.utils import get_column_letter
    for index, spec in matching_scales(headers):
        column = get_column_letter(index + 1)
        scale = ColorScale(
            cfvo=[FormatObject(type=t['type'], val=t.get('value')) for t in spec['thresholds']],
            color=[Color(rgb=color.lstrip('#')) for color in spec['colors']],
        )
        sheet.conditional_formatting.add(
            f'{column}{first_row}:{column}{first_row + row_count - 1}',
            Rule(type='colorScale', colorScale=scale),
        )
    for index, header in enumerate(headers):
        if header == '成本模式':
            column = get_column_letter(index + 1)
            for label, color in COST_COLORS.items():
                sheet.conditional_formatting.add(f'{column}{first_row}:{column}{first_row + row_count - 1}',
                    CellIsRule(operator='equal', formula=['"' + label + '"'],
                               fill=PatternFill('solid', fgColor=color.lstrip('#'))))


def add_artifact_scales(sheet, headers, first_row, row_count):
    if row_count < 1:
        return
    for index, spec in matching_scales(headers):
        sheet.get_range_by_indexes(first_row - 1, index, row_count, 1).conditional_formats.add_color_scale(
            {'colors': spec['colors'], 'thresholds': spec['thresholds']})
    for index, header in enumerate(headers):
        if header == '成本模式':
            for label, color in COST_COLORS.items():
                sheet.get_range_by_indexes(first_row - 1,index,row_count,1).conditional_formats.add_cell_is(
                    {'operator':'equal','formula':'"'+label+'"','format':{'fill':color}})
